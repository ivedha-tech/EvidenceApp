from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import Select
from selenium.common.exceptions import TimeoutException, NoSuchElementException
import time
import os
from datetime import datetime
from docx import Document
from docx.shared import Inches
import logging
from abc import ABC, abstractmethod
from PIL import Image, ImageDraw, ImageFont
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from office365.runtime.auth.authentication_context import AuthenticationContext
from office365.sharepoint.client_context import ClientContext
from office365.sharepoint.files.file import File

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class PageHandler(ABC):
    """Abstract base class for handling different page types"""
    
    @abstractmethod
    def get_url(self, asn):
        """Return the URL to visit for this page type"""
        pass
    
    @abstractmethod
    def get_page_name(self):
        """Return a descriptive name for this page type"""
        pass
    
    @abstractmethod
    def perform_interactions(self, driver, wait):
        """Perform any required interactions on the page before screenshot"""
        pass

class GitHubProfileHandler(PageHandler):
    """Handler for GitHub profile pages"""
    
    def get_url(self, asn):
        return f'https://github.com/{asn}'
    
    def get_page_name(self):
        return 'GitHub Profile'
    
    def perform_interactions(self, driver, wait):
        """Basic GitHub profile - no interactions needed"""
        try:
            # Wait for profile to load
            wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, ".avatar-user, .Avatar--large, [data-testid='avatar']")))
        except TimeoutException:
            logger.warning("GitHub profile elements not found")

class GitHubRepositoriesHandler(PageHandler):
    """Handler for GitHub repositories page"""
    
    def get_url(self, asn):
        return f'https://github.com/{asn}?tab=repositories'
    
    def get_page_name(self):
        return 'GitHub Repositories'
    
    def perform_interactions(self, driver, wait):
        """Interact with repositories page"""
        try:
            # Wait for repositories tab to be active
            time.sleep(2)
            # Try to click on a language filter if available
            try:
                language_filter = driver.find_element(By.CSS_SELECTOR, "[data-testid='filter-by-language'], .select-menu-button")
                language_filter.click()
                time.sleep(1)
            except NoSuchElementException:
                logger.info("No language filter found")
        except Exception as e:
            logger.warning(f"Error interacting with repositories page: {e}")

class GitHubFollowersHandler(PageHandler):
    """Handler for GitHub followers page"""
    
    def get_url(self, asn):
        return f'https://github.com/{asn}?tab=followers'
    
    def get_page_name(self):
        return 'GitHub Followers'
    
    def perform_interactions(self, driver, wait):
        """Basic followers page - no interactions needed"""
        try:
            time.sleep(2)  # Wait for followers to load
        except Exception as e:
            logger.warning(f"Error on followers page: {e}")

class KibanaHandler(PageHandler):
    """Handler for Kibana login and dashboard"""
    
    def __init__(self, kibana_url="https://your-kibana-instance.com", username="admin", password="password"):
        self.kibana_url = kibana_url
        self.username = username
        self.password = password
    
    def get_url(self, asn):
        # For Kibana, we don't use ASN in URL, just return the base Kibana URL
        return f'{self.kibana_url}/login'
    
    def get_page_name(self):
        return 'Kibana Dashboard'
    
    def perform_interactions(self, driver, wait):
        """Login to Kibana and wait for main dashboard"""
        try:
            logger.info(f"Attempting to login to Kibana at {self.kibana_url}")
            
            # Wait for login page to load
            time.sleep(3)
            
            # Try different possible username field selectors
            username_field = None
            username_selectors = [
                'input[name="username"]',
                'input[name="user"]',
                'input[id="username"]',
                'input[type="text"]',
                'input[placeholder*="username" i]',
                'input[placeholder*="user" i]',
                '.form-control[type="text"]'
            ]
            
            for selector in username_selectors:
                try:
                    username_field = wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, selector)))
                    logger.info(f"Found username field with selector: {selector}")
                    break
                except TimeoutException:
                    continue
            
            if not username_field:
                logger.error("Could not find username field")
                return
            
            # Try different possible password field selectors
            password_field = None
            password_selectors = [
                'input[name="password"]',
                'input[type="password"]',
                'input[id="password"]',
                '.form-control[type="password"]'
            ]
            
            for selector in password_selectors:
                try:
                    password_field = driver.find_element(By.CSS_SELECTOR, selector)
                    logger.info(f"Found password field with selector: {selector}")
                    break
                except NoSuchElementException:
                    continue
            
            if not password_field:
                logger.error("Could not find password field")
                return
            
            # Clear and enter credentials
            username_field.clear()
            username_field.send_keys(self.username)
            logger.info("Entered username")
            
            password_field.clear()
            password_field.send_keys(self.password)
            logger.info("Entered password")
            
            # Find and click login button
            login_button = None
            login_selectors = [
                'button[type="submit"]',
                'input[type="submit"]',
                'button[name="login"]',
                'button[id="login"]',
                '.btn-primary',
                '.login-button',
                'button:contains("Log in")',
                'button:contains("Login")',
                'button:contains("Sign in")'
            ]
            
            for selector in login_selectors:
                try:
                    if ':contains' in selector:
                        # Use XPath for text-based selectors
                        # Extract text from :contains("text") format
                        text_part = selector.split(':contains("')[1].split('")')[0]
                        xpath_selector = f"//button[contains(text(), '{text_part}')]"
                        login_button = driver.find_element(By.XPATH, xpath_selector)
                    else:
                        login_button = driver.find_element(By.CSS_SELECTOR, selector)
                    logger.info(f"Found login button with selector: {selector}")
                    break
                except NoSuchElementException:
                    continue
            
            if not login_button:
                logger.error("Could not find login button")
                return
            
            # Click login button
            login_button.click()
            logger.info("Clicked login button")
            
            # Wait for successful login - look for dashboard elements
            logger.info("Waiting for Kibana dashboard to load...")
            
            # Wait for login page to disappear and dashboard elements to appear
            dashboard_selectors = [
                '[data-test-subj="kibana-logo"]',
                '.globalHeader',
                '.application',
                '.kibana-body',
                '[data-test-subj="discover"]',
                '.euiHeader',
                '.kbnGlobalNav',
                '.chromeNavigation'
            ]
            
            dashboard_loaded = False
            for selector in dashboard_selectors:
                try:
                    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, selector)))
                    logger.info(f"Dashboard loaded - found element: {selector}")
                    dashboard_loaded = True
                    break
                except TimeoutException:
                    continue
            
            if not dashboard_loaded:
                # Alternative check - wait for URL change indicating successful login
                try:
                    wait.until(lambda driver: "/login" not in driver.current_url.lower())
                    logger.info("Login successful - URL changed from login page")
                    dashboard_loaded = True
                except TimeoutException:
                    logger.warning("Could not confirm dashboard load, but proceeding")
            
            # Additional wait for dashboard to fully load
            time.sleep(5)
            logger.info("Kibana login and dashboard load completed")
            
        except Exception as e:
            logger.error(f"Error during Kibana login: {str(e)}")

class ASNEvidenceCollector:
    def __init__(self):
        self.setup_driver()
        self.create_output_directory()
        self.page_handlers = self.get_default_page_handlers()
        self.github_logged_in = False
        self.github_username = None
        self.github_password = None
        
        # SharePoint configuration
        self.sharepoint_url = None
        self.sharepoint_folder = None
        self.sharepoint_username = None
        self.sharepoint_password = None
        
        # Email configuration
        self.smtp_server = None
        self.smtp_port = None
        self.email_username = None
        self.email_password = None
        self.email_recipients = []

    def setup_driver(self):
        """Set up the Chrome WebDriver with appropriate options"""
        chrome_options = Options()
        chrome_options.add_argument('--start-maximized')
        chrome_options.add_argument('--disable-notifications')
        chrome_options.add_argument('--disable-blink-features=AutomationControlled')
        chrome_options.add_experimental_option("excludeSwitches", ["enable-automation"])
        chrome_options.add_experimental_option('useAutomationExtension', False)
        
        # Initialize the Chrome WebDriver
        self.driver = webdriver.Chrome(options=chrome_options)
        self.driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
        self.wait = WebDriverWait(self.driver, 15)

    def set_github_credentials(self, username, password):
        """Set GitHub credentials for authentication"""
        self.github_username = username
        self.github_password = password
        logger.info("GitHub credentials set")
    
    def set_sharepoint_config(self, sharepoint_url, folder_path, username, password):
        """Set SharePoint configuration for document upload"""
        self.sharepoint_url = sharepoint_url
        self.sharepoint_folder = folder_path
        self.sharepoint_username = username
        self.sharepoint_password = password
        logger.info("SharePoint configuration set")
    
    def set_email_config(self, smtp_server, smtp_port, username, password, recipients):
        """Set email configuration for sending reports"""
        self.smtp_server = smtp_server
        self.smtp_port = smtp_port
        self.email_username = username
        self.email_password = password
        self.email_recipients = recipients if isinstance(recipients, list) else [recipients]
        logger.info(f"Email configuration set for {len(self.email_recipients)} recipient(s)")

    def login_to_github(self):
        """Login to GitHub using provided credentials"""
        if not self.github_username or not self.github_password:
            logger.warning("GitHub credentials not provided, skipping login")
            return False

        if self.github_logged_in:
            logger.info("Already logged in to GitHub")
            return True

        try:
            logger.info("Attempting to login to GitHub...")
            self.driver.get("https://github.com/login")
            
            # Wait for login page to load
            username_field = self.wait.until(EC.presence_of_element_located((By.ID, "login_field")))
            password_field = self.driver.find_element(By.ID, "password")
            
            # Enter credentials
            username_field.clear()
            username_field.send_keys(self.github_username)
            password_field.clear()
            password_field.send_keys(self.github_password)
            
            # Click login button
            login_button = self.driver.find_element(By.NAME, "commit")
            login_button.click()
            
            # Wait for login to complete - check for successful login
            try:
                # Wait for either dashboard or 2FA page
                self.wait.until(lambda driver: 
                    "github.com/login" not in driver.current_url or
                    driver.find_elements(By.CSS_SELECTOR, "[data-target='sessions.webauthn-challenge']") or
                    driver.find_elements(By.CSS_SELECTOR, "input[name='otp']")
                )
                
                # Check if we're on 2FA page
                if self.driver.find_elements(By.CSS_SELECTOR, "input[name='otp']"):
                    logger.info("Two-factor authentication required")
                    print("\n" + "="*50)
                    print("TWO-FACTOR AUTHENTICATION REQUIRED")
                    print("="*50)
                    print("Please check your authenticator app or SMS for the 6-digit code.")
                    
                    otp_code = input("Enter your 6-digit authentication code: ").strip()
                    
                    if otp_code and len(otp_code) == 6:
                        otp_field = self.driver.find_element(By.CSS_SELECTOR, "input[name='otp']")
                        otp_field.clear()
                        otp_field.send_keys(otp_code)
                        
                        # Click verify button
                        verify_button = self.driver.find_element(By.CSS_SELECTOR, "button[type='submit']")
                        verify_button.click()
                        
                        # Wait for 2FA completion
                        self.wait.until(lambda driver: "github.com/login" not in driver.current_url)
                    else:
                        logger.error("Invalid 2FA code provided")
                        return False
                
                # Verify successful login
                time.sleep(3)
                if "github.com/login" not in self.driver.current_url:
                    self.github_logged_in = True
                    logger.info("Successfully logged in to GitHub")
                    return True
                else:
                    logger.error("Login failed - still on login page")
                    return False
                    
            except TimeoutException:
                logger.error("Login timeout or failed")
                return False
                
        except Exception as e:
            logger.error(f"Error during GitHub login: {str(e)}")
            return False

    def ensure_github_login(self):
        """Ensure GitHub login before accessing GitHub pages"""
        if not self.github_logged_in and (self.github_username and self.github_password):
            return self.login_to_github()
        return self.github_logged_in or not (self.github_username and self.github_password)

    def create_output_directory(self):
        """Create a timestamped directory for storing evidence"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        self.output_dir = f'ASN_Evidence_{timestamp}'
        os.makedirs(self.output_dir, exist_ok=True)
        logger.info(f'Created output directory: {self.output_dir}')

    def get_default_page_handlers(self):
        """Return the default list of page handlers"""
        return [
            GitHubProfileHandler(),
            GitHubRepositoriesHandler(),
            GitHubFollowersHandler()
        ]

    def add_page_handler(self, handler):
        """Add a custom page handler"""
        if isinstance(handler, PageHandler):
            self.page_handlers.append(handler)
        else:
            raise ValueError("Handler must inherit from PageHandler")

    def wait_for_page_load(self):
        """Wait for the page to load completely"""
        try:
            self.wait.until(lambda driver: driver.execute_script('return document.readyState') == 'complete')
            time.sleep(3)  # Additional wait for dynamic content
        except TimeoutException:
            logger.warning('Page load timeout')

    def add_timestamp_overlay(self, image_path):
        """Add timestamp overlay to screenshot"""
        try:
            # Open the image
            with Image.open(image_path) as img:
                # Create a drawing context
                draw = ImageDraw.Draw(img)
                
                # Get current timestamp
                timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                
                # Try to use a default system font, fallback to basic if not available
                try:
                    # Try different font sizes and paths
                    font_size = 32  # Increased from 16 to make it 2x bigger and more readable
                    font_paths = [
                        "arial.ttf",  # Windows
                        "/System/Library/Fonts/Arial.ttf",  # macOS
                        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",  # Linux
                        "C:/Windows/Fonts/arial.ttf"  # Windows absolute path
                    ]
                    
                    font = None
                    for font_path in font_paths:
                        try:
                            font = ImageFont.truetype(font_path, font_size)
                            break
                        except (OSError, IOError):
                            continue
                    
                    if font is None:
                        font = ImageFont.load_default()
                        
                except Exception:
                    font = ImageFont.load_default()
                
                # Calculate text dimensions
                bbox = draw.textbbox((0, 0), timestamp, font=font)
                text_width = bbox[2] - bbox[0]
                text_height = bbox[3] - bbox[1]
                
                # Calculate position (bottom right with padding)
                img_width, img_height = img.size
                padding = 10
                x = img_width - text_width - padding
                y = img_height - text_height - padding
                
                # Create background rectangle (red background)
                bg_padding = 5
                bg_x1 = x - bg_padding
                bg_y1 = y - bg_padding
                bg_x2 = x + text_width + bg_padding
                bg_y2 = y + text_height + bg_padding
                
                # Draw red background rectangle
                draw.rectangle([bg_x1, bg_y1, bg_x2, bg_y2], fill='red')
                
                # Draw white text on top
                draw.text((x, y), timestamp, font=font, fill='white')
                
                # Save the modified image
                img.save(image_path)
                logger.info(f'Timestamp overlay added to: {image_path}')
                
        except Exception as e:
            logger.error(f'Error adding timestamp overlay: {str(e)}')

    def capture_screenshot(self, asn, page_name, page_index):
        """Capture screenshot of the current page"""
        try:
            self.wait_for_page_load()
            screenshot_filename = f'{asn}_{page_index:02d}_{page_name.replace(" ", "_")}.png'
            screenshot_path = os.path.join(self.output_dir, screenshot_filename)
            self.driver.save_screenshot(screenshot_path)
            
            # Add timestamp overlay to the screenshot
            self.add_timestamp_overlay(screenshot_path)
            
            logger.info(f'Screenshot saved with timestamp: {screenshot_path}')
            return screenshot_path
        except Exception as e:
            logger.error(f'Error capturing screenshot: {str(e)}')
            return None

    def process_page(self, asn, handler, page_index):
        """Process a single page for an ASN"""
        logger.info(f'Processing {handler.get_page_name()} for ASN: {asn}')
        
        try:
            # Check if this is a GitHub page and ensure login
            url = handler.get_url(asn)
            if "github.com" in url and not self.ensure_github_login():
                logger.error("Failed to login to GitHub, skipping GitHub pages")
                return {
                    'page_name': handler.get_page_name(),
                    'url': url,
                    'screenshot_path': None,
                    'success': False,
                    'error': 'GitHub login failed'
                }
            
            self.driver.get(url)
            
            # Perform page-specific interactions
            handler.perform_interactions(self.driver, self.wait)
            
            # Capture screenshot
            screenshot_path = self.capture_screenshot(asn, handler.get_page_name(), page_index)
            
            return {
                'page_name': handler.get_page_name(),
                'url': url,
                'screenshot_path': screenshot_path,
                'success': screenshot_path is not None
            }
            
        except Exception as e:
            logger.error(f'Error processing {handler.get_page_name()} for ASN {asn}: {str(e)}')
            return {
                'page_name': handler.get_page_name(),
                'url': handler.get_url(asn),
                'screenshot_path': None,
                'success': False,
                'error': str(e)
            }

    def create_asn_document(self, asn, page_results):
        """Create a Word document for a single ASN with all its screenshots"""
        doc = Document()
        doc.add_heading(f'ASN Evidence Report: {asn}', 0)
        doc.add_paragraph(f'Generated on: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        doc.add_paragraph('')
        
        # Add summary table
        doc.add_heading('Summary', level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Light Grid Accent 1'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Page'
        hdr_cells[1].text = 'Status'
        hdr_cells[2].text = 'URL'
        
        for result in page_results:
            row_cells = table.add_row().cells
            row_cells[0].text = result['page_name']
            row_cells[1].text = 'Success' if result['success'] else 'Failed'
            row_cells[2].text = result['url']
        
        doc.add_paragraph('')
        
        # Add screenshots
        doc.add_heading('Screenshots', level=1)
        
        for i, result in enumerate(page_results, 1):
            if result['success'] and result['screenshot_path'] and os.path.exists(result['screenshot_path']):
                doc.add_heading(f'{i}. {result["page_name"]}', level=2)
                doc.add_paragraph(f'URL: {result["url"]}')
                doc.add_picture(result['screenshot_path'], width=Inches(6.5))
                if i < len(page_results):  # Don't add page break after last screenshot
                    doc.add_page_break()
            else:
                doc.add_heading(f'{i}. {result["page_name"]} (Failed)', level=2)
                doc.add_paragraph(f'URL: {result["url"]}')
                error_msg = result.get('error', 'Unknown error occurred')
                doc.add_paragraph(f'Error: {error_msg}')
                doc.add_paragraph('')

        # Save document
        doc_filename = f'{asn}_Evidence_Report.docx'
        doc_path = os.path.join(self.output_dir, doc_filename)
        doc.save(doc_path)
        logger.info(f'ASN document created: {doc_path}')
        return doc_path

    def process_asn(self, asn):
        """Process a single ASN by visiting all configured pages"""
        logger.info(f'Starting processing for ASN: {asn}')
        page_results = []
        
        try:
            for i, handler in enumerate(self.page_handlers, 1):
                result = self.process_page(asn, handler, i)
                page_results.append(result)
                
                # Add delay between pages to avoid being rate-limited
                time.sleep(2)
            
            # Create Word document for this ASN
            doc_path = self.create_asn_document(asn, page_results)
            
            return {
                'asn': asn,
                'pages_processed': len(page_results),
                'successful_pages': sum(1 for r in page_results if r['success']),
                'document_path': doc_path,
                'page_results': page_results
            }
            
        except Exception as e:
            logger.error(f'Error processing ASN {asn}: {str(e)}')
            return {
                'asn': asn,
                'pages_processed': 0,
                'successful_pages': 0,
                'document_path': None,
                'error': str(e)
            }

    def process_asn_list(self, asn_list):
        """Process a list of ASNs"""
        results = []
        
        try:
            for asn in asn_list:
                asn = asn.strip()
                if asn:
                    result = self.process_asn(asn)
                    results.append(result)
                    
                    # Add delay between ASNs
                    time.sleep(3)

            # Print summary
            self.print_summary(results)
            
        except Exception as e:
            logger.error(f'Error in processing ASN list: {str(e)}')
        finally:
            self.cleanup()

    def print_summary(self, results):
        """Print a summary of processing results"""
        print("\n" + "="*50)
        print("PROCESSING SUMMARY")
        print("="*50)
        
        for result in results:
            print(f"\nASN: {result['asn']}")
            if 'error' in result:
                print(f"  Status: FAILED - {result['error']}")
            else:
                print(f"  Pages Processed: {result['pages_processed']}")
                print(f"  Successful Screenshots: {result['successful_pages']}")
                print(f"  Document: {result['document_path']}")

    def cleanup(self):
        """Clean up resources"""
        try:
            self.driver.quit()
            logger.info('WebDriver closed successfully')
        except Exception as e:
            logger.error(f'Error during cleanup: {str(e)}')

# Example of how to create a custom page handler
class CustomPageHandler(PageHandler):
    """Template for creating custom page handlers"""
    
    def __init__(self, url_template, page_name):
        self.url_template = url_template
        self.page_name = page_name
    
    def get_url(self, asn):
        return self.url_template.format(asn=asn)
    
    def get_page_name(self):
        return self.page_name
    
    def perform_interactions(self, driver, wait):
        """Override this method to add custom interactions"""
        try:
            # Example interactions:
            # 1. Click a button
            # button = wait.until(EC.element_to_be_clickable((By.ID, "some-button")))
            # button.click()
            
            # 2. Select from dropdown
            # dropdown = Select(driver.find_element(By.ID, "some-dropdown"))
            # dropdown.select_by_visible_text("Option")
            
            # 3. Fill form field
            # input_field = driver.find_element(By.NAME, "search")
            # input_field.send_keys("search term")
            
            time.sleep(2)  # Wait for interactions to complete
        except Exception as e:
            logger.warning(f"Error in custom interactions: {e}")

def main():
    print("ASN Evidence Collector")
    print("=====================")
    
    # Get ASN list from user
    asn_input = input("Enter ASN numbers (comma-separated): ")
    asn_list = [asn.strip() for asn in asn_input.split(',') if asn.strip()]
    
    if not asn_list:
        print("No valid ASNs provided. Exiting...")
        return
    
    # Ask for GitHub credentials
    print("\n" + "="*50)
    print("GITHUB AUTHENTICATION")
    print("="*50)
    print("GitHub credentials are required to access profiles and repositories.")
    print("Your credentials will only be used for this session and not stored.")
    
    use_login = input("\nDo you want to login to GitHub? (y/n): ").lower().strip()
    
    github_username = None
    github_password = None
    
    if use_login in ['y', 'yes']:
        github_username = input("GitHub username/email: ").strip()
        if github_username:
            import getpass
            github_password = getpass.getpass("GitHub password: ")
            if not github_password:
                print("Warning: No password provided. Some GitHub pages may not be accessible.")
        else:
            print("Warning: No username provided. Some GitHub pages may not be accessible.")
    else:
        print("Warning: Skipping GitHub login. Some pages may not be accessible without authentication.")
    
    # Ask for Kibana configuration
    print("\n" + "="*50)
    print("KIBANA CONFIGURATION")
    print("="*50)
    use_kibana = input("Do you want to include Kibana dashboard? (y/n): ").lower().strip()
    
    print(f"\nProcessing {len(asn_list)} ASN(s)...")
    base_handlers = len(ASNEvidenceCollector().get_default_page_handlers())
    total_handlers = base_handlers + (1 if use_kibana in ['y', 'yes'] else 0)
    print(f"Each ASN will be processed across {total_handlers} pages")
    
    # Initialize and run the collector
    collector = ASNEvidenceCollector()
    
    # Set GitHub credentials if provided
    if github_username and github_password:
        collector.set_github_credentials(github_username, github_password)
    
    # Add Kibana handler if requested
    if use_kibana in ['y', 'yes']:
        # Hardcoded Kibana credentials - modify these for your Kibana instance
        kibana_url = "https://your-kibana-instance.com"  # Change this to your Kibana URL
        kibana_username = "admin"  # Change this to your Kibana username
        kibana_password = "password"  # Change this to your Kibana password
        
        kibana_handler = KibanaHandler(kibana_url, kibana_username, kibana_password)
        collector.add_page_handler(kibana_handler)
        print(f"Added Kibana handler for: {kibana_url}")
    
    # Example of adding other custom page handlers:
    # custom_handler = CustomPageHandler("https://example.com/{asn}", "Custom Page")
    # collector.add_page_handler(custom_handler)
    
    collector.process_asn_list(asn_list)
    
    print(f"\nProcessing completed!")
    print(f"Evidence documents have been saved in: {collector.output_dir}")

if __name__ == "__main__":
    main() 